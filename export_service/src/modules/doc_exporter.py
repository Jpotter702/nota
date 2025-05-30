from docx import Document

def export(input_path: str, output_path: str):
    doc = Document()

    with open(input_path, "r", encoding="utf-8") as f:
        for line in f:
            if line.strip().startswith("# "):
                doc.add_heading(line.strip("# ").strip(), level=1)
            elif line.strip().startswith("## "):
                doc.add_heading(line.strip("# ").strip(), level=2)
            elif line.strip():
                doc.add_paragraph(line.strip())

    doc.save(output_path)
    print(f"📄 DOCX exported to {output_path}")